import io
import os
import re
import zipfile
import tempfile
import traceback
from dataclasses import dataclass, field
from typing import List, Tuple, Optional, Dict, Any

import numpy as np
import pandas as pd
import streamlit as st

import pdfplumber
from pdf2image import convert_from_bytes
import pytesseract
import cv2

from rapidfuzz import fuzz

# Optional openpyxl imports for formatting
try:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    pass

# camelot is optional at import time
try:
    import camelot
    CAMELOT_AVAILABLE = True
except Exception:
    CAMELOT_AVAILABLE = False


# ==================================================================================
# CONFIGURATION
# ==================================================================================

st.set_page_config(
    page_title="Multi-PDF to Excel Converter",
    page_icon="📊",
    layout="wide",
)

TEXT_DENSITY_THRESHOLD = 40
HEADER_SIMILARITY_THRESHOLD = 78   
COLUMN_COUNT_TOLERANCE = 0         
DTYPE_SIMILARITY_THRESHOLD = 0.6   
EXCEL_SHEET_NAME_MAX_LEN = 31  


# ==================================================================================
# DATA STRUCTURES
# ==================================================================================

@dataclass
class ExtractedTable:
    """A single raw table pulled from a PDF, tagged with provenance."""
    dataframe: pd.DataFrame
    page_number: int
    source: str


@dataclass
class FileResult:
    """Holds the outcome of processing a single uploaded PDF."""
    filename: str
    pdf_type: str = "unknown"
    status: str = "pending"
    error_message: str = ""
    sheet_names: List[str] = field(default_factory=list)
    preview_frames: Dict[str, pd.DataFrame] = field(default_factory=dict)
    excel_bytes: Optional[bytes] = None
    tables: List[pd.DataFrame] = field(default_factory=list)
    num_pages: int = 0
    num_tables_found: int = 0
    num_tables_after_merge: int = 0


# ==================================================================================
# STEP 1: PDF TYPE DETECTION
# ==================================================================================

def detect_pdf_type(pdf_bytes: bytes) -> Tuple[str, int]:
    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            num_pages = len(pdf.pages)
            if num_pages == 0:
                return "scanned", 0

            total_chars = 0
            pages_sampled = min(num_pages, 5)
            for page in pdf.pages[:pages_sampled]:
                text = page.extract_text() or ""
                total_chars += len(text.strip())

            avg_chars_per_page = total_chars / max(pages_sampled, 1)

            if avg_chars_per_page >= TEXT_DENSITY_THRESHOLD:
                return "text", num_pages
            else:
                return "scanned", num_pages
    except Exception:
        return "scanned", 0


# ==================================================================================
# STEP 2A: TEXT-BASED EXTRACTION
# ==================================================================================

def extract_text_pdf(pdf_bytes: bytes, filename: str) -> List[ExtractedTable]:
    extracted: List[ExtractedTable] = []
    camelot_pages_with_tables = set()

    if CAMELOT_AVAILABLE:
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
                tmp.write(pdf_bytes)
                tmp_path = tmp.name

            for flavor in ("lattice", "stream"):
                try:
                    tables = camelot.read_pdf(tmp_path, pages="all", flavor=flavor)
                except Exception:
                    tables = []
                for t in tables:
                    df = t.df
                    if df is not None and not df.empty:
                        page_num = int(t.page) if hasattr(t, "page") else 0
                        extracted.append(ExtractedTable(df, page_num, f"camelot-{flavor}"))
                        camelot_pages_with_tables.add(page_num)
                if extracted:
                    break
        except Exception:
            pass
        finally:
            if tmp_path and os.path.exists(tmp_path):
                try:
                    os.remove(tmp_path)
                except OSError:
                    pass

    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                if i in camelot_pages_with_tables:
                    continue
                try:
                    page_tables = page.extract_tables()
                except Exception:
                    page_tables = []

                if page_tables:
                    for raw_table in page_tables:
                        if not raw_table:
                            continue
                        df = pd.DataFrame(raw_table)
                        if not df.empty:
                            extracted.append(ExtractedTable(df, i, "pdfplumber"))
                else:
                    text = page.extract_text() or ""
                    rows = []
                    for line in text.split("\n"):
                        line = line.strip()
                        if not line:
                            continue
                        cols = re.split(r"\s{2,}|\t", line)
                        rows.append(cols)
                    if rows:
                        max_cols = max(len(r) for r in rows)
                        rows = [r + [""] * (max_cols - len(r)) for r in rows]
                        df = pd.DataFrame(rows)
                        if not df.empty:
                            extracted.append(ExtractedTable(df, i, "pdfplumber-text"))
    except Exception:
        pass

    return extracted


# ==================================================================================
# STEP 2B: SCANNED PDF EXTRACTION (OCR)
# ==================================================================================

def preprocess_image(image: np.ndarray) -> np.ndarray:
    gray = cv2.cvtColor(image, cv2.COLOR_RGB2GRAY)
    denoised = cv2.fastNlMeansDenoising(gray, h=10)
    thresh = cv2.adaptiveThreshold(
        denoised, 255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY,
        31, 15
    )
    kernel = np.ones((1, 1), np.uint8)
    processed = cv2.morphologyEx(thresh, cv2.MORPH_CLOSE, kernel)
    return processed

def _extract_annexure_template(ocr_data: Dict[str, list], img_width: int) -> List[pd.DataFrame]:
    """
    Template-aware parser specifically for 'Annexure to Supplementary Invoice'.
    Isolates Header, Pricing Table, and Invoice Table into exact structured dataframes.
    Uses horizontal constraints, Regex stitching, and 1D K-Means for column snapping.
    """
    # 1. Parse words
    words = []
    n = len(ocr_data["text"])
    for i in range(n):
        text = str(ocr_data["text"][i]).strip()
        if not text: continue
        try:
            conf = float(ocr_data.get("conf", ["0"] * n)[i])
            if conf < 15: continue
        except: continue
        
        x, y, w, h = ocr_data["left"][i], ocr_data["top"][i], ocr_data["width"][i], ocr_data["height"][i]
        words.append({
            "text": text, "left": x, "top": y, "right": x+w, "bottom": y+h,
            "width": w, "height": h, "cx": x + w/2, "cy": y + h/2
        })
        
    if not words: return []
    
    # 2. Fix Broken Numbers (e.g., "2,110" and ".31" split by a gap)
    words.sort(key=lambda w: (w["cy"], w["left"]))
    fixed_words = []
    skip = False
    for i in range(len(words)-1):
        if skip:
            skip = False
            continue
        w1, w2 = words[i], words[i+1]
        
        # If words are on the same line and very close horizontally
        if abs(w1["cy"] - w2["cy"]) < 15 and (w2["left"] - w1["right"]) < 30:
            # Check if one ends with digit/comma and the next starts with dot/digit
            if (re.search(r'[\d\,]$', w1["text"]) and re.search(r'^[\.\,]\s*\d', w2["text"])) or \
               (w1["text"].isdigit() and w2["text"].isdigit()):
                w_new = w1.copy()
                w_new["text"] = w1["text"] + w2["text"].replace(" ", "")
                w_new["right"] = w2["right"]
                w_new["width"] = w2["right"] - w1["left"]
                w_new["cx"] = (w_new["left"] + w_new["right"]) / 2
                fixed_words.append(w_new)
                skip = True
                continue
        fixed_words.append(w1)
    if not skip and words: fixed_words.append(words[-1])
    words = fixed_words
    
    # 3. Split merged numeric/text like "1,788.40J3" -> "1,788.40" and "J3"
    split_words = []
    for w in words:
        m = re.match(r'^([\d\,\.]+)([a-zA-Z0-9]{2,3})$', w["text"])
        # If it matched the pattern and the first part actually has digits (not just dots)
        if m and any(c.isdigit() for c in m.group(1)) and not w["text"].isdigit():
            w1 = w.copy()
            w1["text"] = m.group(1)
            # Estimate geometry based on length ratios
            ratio = len(w1["text"]) / len(w["text"])
            w1["right"] = w["left"] + int(w["width"] * ratio)
            w1["width"] = w1["right"] - w1["left"]
            w1["cx"] = w1["left"] + w1["width"] / 2
            
            w2 = w.copy()
            w2["text"] = m.group(2)
            w2["left"] = w1["right"]
            w2["width"] = w["right"] - w2["left"]
            w2["cx"] = w2["left"] + w2["width"] / 2
            split_words.extend([w1, w2])
        else:
            split_words.append(w)
    words = split_words

    # 4. Group into horizontal lines
    words.sort(key=lambda w: w["cy"])
    lines = []
    curr_line = [words[0]]
    for w in words[1:]:
        avg_cy = sum(x["cy"] for x in curr_line) / len(curr_line)
        avg_h = sum(x["height"] for x in curr_line) / len(curr_line)
        if abs(w["cy"] - avg_cy) < max(avg_h * 0.5, 8):
            curr_line.append(w)
        else:
            curr_line.sort(key=lambda x: x["left"])
            lines.append(curr_line)
            curr_line = [w]
    if curr_line:
        curr_line.sort(key=lambda x: x["left"])
        lines.append(curr_line)

    dfs = []
    
    # 5. Extract Header Dictionary
    header_keys = ["PSF No.", "SA No.", "Plant", "Material No.", "Annexure Number", "Annexure Date"]
    search_keys = ["PSF No", "SA No", "Plant", "Material No", "Annexure Number", "Annexure Date"]
    header_dict = {}
    
    for line in lines:
        for idx, hk in enumerate(search_keys):
            clean_hk = hk.replace(" ", "").lower()
            for i, w in enumerate(line):
                clean_w = w["text"].replace(".", "").replace(" ", "").lower()
                # Check for match (fuzzy or direct substring)
                if fuzz.ratio(clean_hk, clean_w) > 85 or (len(clean_w)>4 and clean_w in clean_hk):
                    if i + 1 < len(line):
                        next_w = line[i+1]
                        # Check if the next word is another key (means value is empty)
                        is_key = any(fuzz.ratio(k.replace(" ","").lower(), next_w["text"].replace(".","").replace(" ","").lower()) > 85 for k in search_keys)
                        
                        if not is_key and (next_w["left"] - w["right"]) < 350:
                            val_text = next_w["text"]
                            j = i + 2
                            while j < len(line):
                                is_next_key = any(fuzz.ratio(k.replace(" ","").lower(), line[j]["text"].replace(".","").replace(" ","").lower()) > 85 for k in search_keys)
                                if is_next_key or (line[j]["left"] - line[j-1]["right"] > 80):
                                    break
                                val_text += " " + line[j]["text"]
                                j += 1
                            header_dict[header_keys[idx]] = val_text
                            
    if header_dict:
        dfs.append(pd.DataFrame([header_dict]))
        
    # 6. Locate Table Boundaries
    pricing_start, invoice_start = -1, -1
    for i, line in enumerate(lines):
        line_str = " ".join([w["text"].lower() for w in line])
        if "effective" in line_str and "date" in line_str and "price" in line_str:
            pricing_start = i
        if "invoice no" in line_str and "grn" in line_str:
            invoice_start = i
            
    def _build_table(table_lines, expected_cols):
        if not table_lines: return pd.DataFrame()
        all_words = [w for l in table_lines for w in l]
        if not all_words: return pd.DataFrame()
        
        num_cols = len(expected_cols)
        cx_vals = [w["cx"] for w in all_words]
        
        # Initial cluster centers spread evenly across the X-axis bounds
        centers = np.linspace(min(cx_vals), max(cx_vals), num_cols)
        
        # 1D K-Means to find robust column anchor points
        for _ in range(15):
            clusters = [[] for _ in range(num_cols)]
            for cx in cx_vals:
                best_idx = int(np.argmin([abs(cx - c) for c in centers]))
                clusters[best_idx].append(cx)
            new_centers = []
            for idx, cl in enumerate(clusters):
                new_centers.append(np.mean(cl) if cl else centers[idx])
            centers = new_centers
            
        centers = sorted(centers)
        data = []
        
        for line in table_lines:
            line_str = " ".join([w["text"].lower() for w in line])
            # Skip rows consisting mostly of header text (avoid repeating headers inside data)
            if sum(1 for c in expected_cols if fuzz.ratio(c.lower(), line_str) > 60) > 3:
                continue
            # Skip sparse rows unless it's a totals row
            if len(line) < 2 and "total" not in line_str:
                continue
                
            row_data = [""] * num_cols
            for w in line:
                best_col = int(np.argmin([abs(w["cx"] - c) for c in centers]))
                if row_data[best_col]:
                    row_data[best_col] += " " + w["text"]
                else:
                    row_data[best_col] = w["text"]
            
            str_row = " ".join(row_data).lower()
            if "invoice no" in str_row or "effective from" in str_row: continue
            data.append(row_data)
            
        return pd.DataFrame(data, columns=expected_cols)
        
    # 7. Extract Pricing Table
    if pricing_start != -1:
        end_idx = invoice_start if invoice_start != -1 else len(lines)
        pricing_df = _build_table(lines[pricing_start:end_idx], [
            "Effective From Date", "Effective To Date", "Old Base Price", 
            "New Base Price", "Old Freight", "New Freight", 
            "Old Packing", "New Packing", "Net Diff"
        ])
        if not pricing_df.empty: dfs.append(pricing_df)
            
    # 8. Extract Invoice Table
    if invoice_start != -1:
        invoice_df = _build_table(lines[invoice_start:], [
            "Invoice No", "Invoice Date", "GRN No", "Qty", 
            "Original Invoice", "Net Diff", "Net Value", "Tax Code", 
            "SGST", "CGST", "IGST", "Total Value"
        ])
        if not invoice_df.empty: dfs.append(invoice_df)
            
    # 9. Continuation Page Catch (If the page is purely a continued invoice table)
    if pricing_start == -1 and invoice_start == -1 and len(lines) > 5:
        # Determine likely column count
        avg_words_per_line = sum(len(l) for l in lines) / len(lines)
        if avg_words_per_line >= 8:
            invoice_df = _build_table(lines, [
                "Invoice No", "Invoice Date", "GRN No", "Qty", 
                "Original Invoice", "Net Diff", "Net Value", "Tax Code", 
                "SGST", "CGST", "IGST", "Total Value"
            ])
            if not invoice_df.empty: dfs.append(invoice_df)
            
    return dfs

def _words_to_table(ocr_data: Dict[str, list], y_tolerance: int = 12) -> pd.DataFrame:
    words = []
    n = len(ocr_data["text"])
    for i in range(n):
        text = ocr_data["text"][i].strip()
        if not text:
            continue
        conf = ocr_data.get("conf", ["0"] * n)[i]
        try:
            conf_val = float(conf)
        except (ValueError, TypeError):
            conf_val = -1
        if conf_val != -1 and conf_val < 30:
            continue
        words.append({
            "text": text,
            "left": ocr_data["left"][i],
            "top": ocr_data["top"][i],
        })

    if not words:
        return pd.DataFrame()

    words.sort(key=lambda w: w["top"])
    rows = []
    current_row = [words[0]]
    current_top = words[0]["top"]

    for w in words[1:]:
        if abs(w["top"] - current_top) <= y_tolerance:
            current_row.append(w)
        else:
            rows.append(current_row)
            current_row = [w]
            current_top = w["top"]
    rows.append(current_row)

    table_rows = []
    for row in rows:
        row_sorted = sorted(row, key=lambda w: w["left"])
        table_rows.append([w["text"] for w in row_sorted])

    max_cols = max(len(r) for r in table_rows)
    table_rows = [r + [""] * (max_cols - len(r)) for r in table_rows]

    return pd.DataFrame(table_rows)

def extract_ocr_pdf(pdf_bytes: bytes, filename: str, dpi: int = 200) -> List[ExtractedTable]:
    extracted: List[ExtractedTable] = []
    try:
        pages = convert_from_bytes(pdf_bytes, dpi=dpi)
    except Exception as e:
        raise RuntimeError(f"Failed to rasterize PDF for OCR: {e}")

    for i, pil_page in enumerate(pages, start=1):
        try:
            img = np.array(pil_page.convert("RGB"))
            processed = preprocess_image(img)
            ocr_data = pytesseract.image_to_data(
                processed, output_type=pytesseract.Output.DICT
            )
            
            # Smart Routing: Check for template keywords
            text_content = " ".join([str(x).lower() for x in ocr_data.get("text", []) if str(x).strip()])
            is_annexure = (
                "annexure to supplementary invoice" in text_content or 
                ("effective from date" in text_content and "new base price" in text_content) or
                ("invoice no" in text_content and "grn no" in text_content and "tax code" in text_content)
            )
            
            if is_annexure:
                dfs = _extract_annexure_template(ocr_data, img.shape[1])
                for df in dfs:
                    if not df.empty:
                        extracted.append(ExtractedTable(df, i, "ocr-annexure"))
            else:
                df = _words_to_table(ocr_data)
                if not df.empty:
                    extracted.append(ExtractedTable(df, i, "ocr-generic"))
                    
        except Exception:
            continue

    return extracted


# ==================================================================================
# STEP 3: CONTINUITY DETECTION & MERGING
# ==================================================================================

def _infer_column_dtype_signature(series: pd.Series) -> str:
    sample = series.dropna().astype(str).str.strip()
    sample = sample[sample != ""].head(20)
    if sample.empty:
        return "empty"
    numeric_count = sum(bool(re.fullmatch(r"-?\d+(\.\d+)?%?", v)) for v in sample)
    date_count = sum(bool(re.fullmatch(r"\d{1,4}[/-]\d{1,2}[/-]\d{1,4}", v)) for v in sample)
    if numeric_count / len(sample) >= 0.6:
        return "numeric"
    if date_count / len(sample) >= 0.6:
        return "date"
    return "text"

def _dtype_signature_similarity(df1: pd.DataFrame, df2: pd.DataFrame) -> float:
    n = min(len(df1.columns), len(df2.columns))
    if n == 0: return 0.0
    matches = 0
    for c in range(n):
        sig1 = _infer_column_dtype_signature(df1.iloc[:, c])
        sig2 = _infer_column_dtype_signature(df2.iloc[:, c])
        if sig1 == sig2:
            matches += 1
    return matches / n

def _header_similarity(df1: pd.DataFrame, df2: pd.DataFrame) -> float:
    header1 = [str(x) for x in df1.iloc[0].tolist()] if len(df1) else []
    header2 = [str(x) for x in df2.iloc[0].tolist()] if len(df2) else []
    if not header1 or not header2: return 0.0
    n = min(len(header1), len(header2))
    scores = [fuzz.ratio(header1[i], header2[i]) for i in range(n)]
    return sum(scores) / len(scores) if scores else 0.0

def are_tables_similar(df1: pd.DataFrame, df2: pd.DataFrame) -> bool:
    if df1.empty or df2.empty: return False
    col_diff = abs(len(df1.columns) - len(df2.columns))
    if col_diff > COLUMN_COUNT_TOLERANCE: return False
    header_score = _header_similarity(df1, df2)
    dtype_score = _dtype_signature_similarity(df1, df2)

    if header_score >= HEADER_SIMILARITY_THRESHOLD: return True
    if dtype_score >= DTYPE_SIMILARITY_THRESHOLD and header_score >= 40: return True
    return False

def merge_tables(tables: List[pd.DataFrame]) -> pd.DataFrame:
    if not tables: return pd.DataFrame()
    if len(tables) == 1: return tables[0].reset_index(drop=True)

    base_header = [str(x).strip().lower() for x in tables[0].iloc[0].tolist()]
    merged_frames = [tables[0]]

    for t in tables[1:]:
        if t.empty: continue
        first_row = [str(x).strip().lower() for x in t.iloc[0].tolist()]
        n = min(len(first_row), len(base_header))
        if n and fuzz.ratio(" ".join(first_row[:n]), " ".join(base_header[:n])) >= HEADER_SIMILARITY_THRESHOLD:
            merged_frames.append(t.iloc[1:])
        else:
            merged_frames.append(t)

    return pd.concat(merged_frames, ignore_index=True, sort=False)

def group_and_merge_tables(extracted_tables: List[ExtractedTable]) -> List[pd.DataFrame]:
    if not extracted_tables: return []
    ordered = sorted(extracted_tables, key=lambda t: t.page_number)
    groups: List[List[pd.DataFrame]] = [[ordered[0].dataframe]]

    for et in ordered[1:]:
        last_group = groups[-1]
        last_df = last_group[-1]
        if are_tables_similar(last_df, et.dataframe):
            last_group.append(et.dataframe)
        else:
            groups.append([et.dataframe])

    return [merge_tables(g) for g in groups]


# ==================================================================================
# STEP 4: DATA CLEANING
# ==================================================================================

def clean_tables(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty: return pd.DataFrame()
    df = df.copy()

    df = df.astype(str).apply(lambda col: col.str.replace(r"\s+", " ", regex=True).str.strip())
    df = df.replace(r"^\s*$", np.nan, regex=True)
    df = df.replace("nan", np.nan)
    
    # Drop rows/columns that are completely empty
    df = df.dropna(axis=0, how="all")
    df = df.dropna(axis=1, how="all")

    if df.empty: return df
    df = df.reset_index(drop=True)

    # If the dataframe already has valid string columns (like our template output), don't push them down
    if all(isinstance(c, str) and not str(c).startswith("col_") and not str(c).isdigit() for c in df.columns):
        pass # Already formatted well
    else:
        # Standard extraction logic: promote first row to header
        first_row = df.iloc[0]
        non_null_ratio = first_row.notna().mean()
        if non_null_ratio > 0.5:
            df.columns = [str(c) if pd.notna(c) else f"col_{i}" for i, c in enumerate(first_row)]
            df = df.iloc[1:].reset_index(drop=True)
        else:
            df.columns = [f"col_{i}" for i in range(len(df.columns))]

    # Deduplicate column names
    seen = {}
    new_cols = []
    for c in df.columns:
        c = str(c).strip() or "col"
        if c in seen:
            seen[c] += 1
            new_cols.append(f"{c}_{seen[c]}")
        else:
            seen[c] = 0
            new_cols.append(c)
    df.columns = new_cols

    # Remove inline repeated headers
    header_lower = [str(c).strip().lower() for c in df.columns]
    def _is_header_dup(row):
        vals = [str(v).strip().lower() for v in row.tolist()]
        n = min(len(vals), len(header_lower))
        if n == 0: return False
        return fuzz.ratio(" ".join(vals[:n]), " ".join(header_lower[:n])) >= HEADER_SIMILARITY_THRESHOLD

    if len(df) > 0:
        dup_mask = df.apply(_is_header_dup, axis=1)
        df = df[~dup_mask]

    df = df.drop_duplicates(keep="first")
    df = df.dropna(axis=0, how="all")
    df = df.reset_index(drop=True)

    return df


# ==================================================================================
# EXCEL WORKBOOK GENERATION (Single Sheet vs Multi Sheet)
# ==================================================================================

def _safe_sheet_name(name: str, used_names: set) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "_", name).strip() or "Sheet"
    name = name[:EXCEL_SHEET_NAME_MAX_LEN]
    base = name
    counter = 1
    while name in used_names:
        suffix = f"_{counter}"
        name = base[: EXCEL_SHEET_NAME_MAX_LEN - len(suffix)] + suffix
        counter += 1
    used_names.add(name)
    return name

def build_excel_workbook(tables: List[pd.DataFrame], layout_mode: str = "multi") -> Tuple[bytes, List[str], Dict[str, pd.DataFrame]]:
    """
    Write a list of cleaned tables into an Excel workbook.
    Adds openpyxl formatting to bold headers and autofit columns perfectly.
    """
    output = io.BytesIO()
    used_names: set = set()
    sheet_names: List[str] = []
    previews: Dict[str, pd.DataFrame] = {}

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        if not tables:
            pd.DataFrame({"Notice": ["No tables could be extracted."]}).to_excel(
                writer, sheet_name="Sheet1", index=False
            )
            sheet_names.append("Sheet1")
        else:
            if layout_mode == "single":
                sheet_name = _safe_sheet_name("All_Tables", used_names)
                sheet_names.append(sheet_name)
                current_row = 0
                
                for idx, df in enumerate(tables):
                    if df.empty:
                        continue
                    if current_row == 0:
                        previews[sheet_name] = df.head(5)
                        
                    label_df = pd.DataFrame(columns=[f"--- Table {idx+1} ---"])
                    label_df.to_excel(writer, sheet_name=sheet_name, startrow=current_row, index=False)
                    df.to_excel(writer, sheet_name=sheet_name, startrow=current_row + 1, index=False)
                    current_row += len(df) + 4
            else:
                for idx, df in enumerate(tables, start=1):
                    if df.empty:
                        continue
                    sheet_name = _safe_sheet_name(f"Table_{idx}", used_names)
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
                    sheet_names.append(sheet_name)
                    previews[sheet_name] = df.head(5)

        # Apply OpenPyXL formatting (Autofit & Bold)
        try:
            for sheet_name in writer.sheets:
                worksheet = writer.sheets[sheet_name]
                for col in worksheet.columns:
                    max_length = 0
                    column = col[0].column_letter # Get the column name
                    for cell in col:
                        try:
                            if cell.value:
                                max_length = max(max_length, len(str(cell.value)))
                        except:
                            pass
                        
                        # Bold styling for label separators or true top rows
                        if cell.row == 1 or (layout_mode == "single" and str(cell.value).startswith("---")):
                            cell.font = Font(bold=True)
                            
                    adjusted_width = (max_length + 2)
                    worksheet.column_dimensions[column].width = min(adjusted_width, 50) # Cap width at 50
        except Exception as e:
            print(f"Excel formatting skipped: {e}")

    return output.getvalue(), sheet_names, previews


# ==================================================================================
# STEP 5 & 6: PIPELINE EXECUTORS
# ==================================================================================

def process_single_pdf(filename: str, pdf_bytes: bytes, layout_mode: str) -> FileResult:
    result = FileResult(filename=filename, status="processing")

    pdf_type, num_pages = detect_pdf_type(pdf_bytes)
    result.pdf_type = pdf_type
    result.num_pages = num_pages

    if pdf_type == "text":
        raw_tables = extract_text_pdf(pdf_bytes, filename)
    else:
        raw_tables = extract_ocr_pdf(pdf_bytes, filename)

    result.num_tables_found = len(raw_tables)
    merged_tables = group_and_merge_tables(raw_tables)
    
    cleaned_tables = [clean_tables(df) for df in merged_tables]
    cleaned_tables = [df for df in cleaned_tables if not df.empty]
    
    result.tables = cleaned_tables
    result.num_tables_after_merge = len(cleaned_tables)

    excel_bytes, sheet_names, previews = build_excel_workbook(cleaned_tables, layout_mode)
    result.excel_bytes = excel_bytes
    result.sheet_names = sheet_names
    result.preview_frames = previews
    result.status = "done"

    return result


def process_multiple_pdfs(
    uploaded_files: list,
    overall_progress_bar,
    status_container,
    layout_mode: str
) -> List[FileResult]:
    results: List[FileResult] = []
    total = len(uploaded_files)

    for i, uploaded_file in enumerate(uploaded_files, start=1):
        filename = uploaded_file.name
        row = status_container.container()
        row_placeholder = row.empty()
        row_placeholder.info(f"⏳ Processing **{filename}** ({i}/{total})...")

        try:
            pdf_bytes = uploaded_file.getvalue()
            result = process_single_pdf(filename, pdf_bytes, layout_mode)
            results.append(result)
            row_placeholder.success(
                f"✅ **{filename}** — type: `{result.pdf_type}` — "
                f"Generated **{len(result.sheet_names)}** sheet(s) "
                f"({result.num_pages} pages)"
            )
        except Exception as e:
            err_result = FileResult(
                filename=filename,
                status="error",
                error_message=f"{type(e).__name__}: {e}",
            )
            results.append(err_result)
            row_placeholder.error(f"❌ **{filename}** failed: {err_result.error_message}")
            traceback.print_exc()

        overall_progress_bar.progress(i / total, text=f"Overall progress: {i}/{total} files")

    return results

def build_zip_of_excels(results: List[FileResult]) -> bytes:
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for r in results:
            if r.status == "done" and r.excel_bytes:
                base_name = os.path.splitext(r.filename)[0]
                zf.writestr(f"{base_name}.xlsx", r.excel_bytes)
    return zip_buffer.getvalue()


# ==================================================================================
# STREAMLIT UI
# ==================================================================================

def main():
    st.title("📊 Multi-PDF to Excel Converter")
    st.caption(
        "Upload multiple PDFs (text-based or scanned). Each PDF is converted into "
        "an Excel workbook, with multi-page tables intelligently merged. Now featuring "
        "template-aware OCR for *Annexure to Supplementary Invoice* layouts."
    )

    if not CAMELOT_AVAILABLE:
        st.warning(
            "⚠️ `camelot-py` (or its Ghostscript dependency) is not available in this "
            "environment. Text-based extraction will fall back to pdfplumber only."
        )

    with st.sidebar:
        st.header("⚙️ Settings")
        
        st.subheader("Data Export Layout")
        sheet_layout = st.radio(
            "How should tables be formatted inside the Excel file?",
            options=[
                "Multi-Sheet (One table per tab)", 
                "Single-Sheet (All tables stacked vertically)"
            ],
            index=0,
            help="Choose whether to place each extracted table on its own Excel tab, or stack them all on a single tab with spacing."
        )
        layout_mode = "single" if "Single" in sheet_layout else "multi"

        st.markdown("---")
        ocr_dpi = st.slider("OCR rasterization DPI", 100, 300, 200, step=25,
                             help="Higher DPI improves OCR accuracy but is slower.")
        
        st.markdown("---")
        st.markdown("**Tech stack:** pdfplumber, camelot, pytesseract, pdf2image, OpenCV, rapidfuzz, openpyxl")

    uploaded_files = st.file_uploader(
        "Upload PDF files (you can select 100+ at once)",
        type=["pdf"],
        accept_multiple_files=True,
    )

    if uploaded_files:
        st.subheader(f"📁 {len(uploaded_files)} file(s) uploaded")
        with st.expander("View uploaded file list", expanded=False):
            for f in uploaded_files:
                st.write(f"• {f.name}  ({f.size / 1024:.1f} KB)")

        col1, col2 = st.columns([1, 4])
        with col1:
            start_clicked = st.button("🚀 Start Conversion", type="primary")

        if start_clicked:
            global TEXT_DENSITY_THRESHOLD
            st.session_state["ocr_dpi"] = ocr_dpi

            st.subheader("🔄 Processing status")
            overall_progress_bar = st.progress(0, text="Overall progress: 0/0 files")
            status_container = st.container()

            global extract_ocr_pdf
            original_ocr_fn = extract_ocr_pdf

            def _ocr_with_dpi(pdf_bytes, filename, dpi=ocr_dpi):
                return original_ocr_fn(pdf_bytes, filename, dpi=dpi)
            extract_ocr_pdf = _ocr_with_dpi

            with st.spinner("Converting PDFs to Excel... this may take a while for large batches."):
                results = process_multiple_pdfs(uploaded_files, overall_progress_bar, status_container, layout_mode)

            extract_ocr_pdf = original_ocr_fn  
            st.session_state["results"] = results
            st.success("🎉 Batch processing complete!")

    # ------------------------------------------------------------------
    # RESULTS / DOWNLOADS
    # ------------------------------------------------------------------
    results: List[FileResult] = st.session_state.get("results", [])

    if results:
        st.markdown("---")
        st.subheader("📥 Download Results")

        done_results = [r for r in results if r.status == "done"]
        error_results = [r for r in results if r.status == "error"]

        summary_cols = st.columns(3)
        summary_cols[0].metric("Total files", len(results))
        summary_cols[1].metric("Succeeded", len(done_results))
        summary_cols[2].metric("Failed", len(error_results))

        if done_results:
            st.markdown("#### Batch Actions")
            colA, colB = st.columns(2)
            
            # Action 1: ZIP of individual files
            zip_bytes = build_zip_of_excels(done_results)
            colA.download_button(
                label=f"⬇️ Download {len(done_results)} files as ZIP",
                data=zip_bytes,
                file_name="converted_excels.zip",
                mime="application/zip",
                type="primary",
                use_container_width=True
            )
            
            # Action 2: Massive master combined file
            master_tables = []
            for r in done_results:
                for t in r.tables:
                    t_copy = t.copy()
                    if "Source PDF" not in t_copy.columns:
                        t_copy.insert(0, "Source PDF", r.filename)
                    master_tables.append(t_copy)
                    
            if master_tables:
                master_excel_bytes, _, _ = build_excel_workbook(master_tables, layout_mode="single")
                colB.download_button(
                    label="⬇️ Download ONE Master Excel (All PDFs Combined)",
                    data=master_excel_bytes,
                    file_name="Master_Combined_Data.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="secondary",
                    use_container_width=True,
                    help="Combine every table from every PDF into a single Excel file."
                )

        st.markdown("#### Individual files")
        for r in results:
            with st.expander(f"{'✅' if r.status == 'done' else '❌'} {r.filename}"):
                if r.status == "error":
                    st.error(f"Processing failed: {r.error_message}")
                    continue

                meta_cols = st.columns(4)
                meta_cols[0].write(f"**Type detected:** `{r.pdf_type}`")
                meta_cols[1].write(f"**Pages:** {r.num_pages}")
                meta_cols[2].write(f"**Tables found:** {r.num_tables_found}")
                meta_cols[3].write(f"**Sheets/Tables post-merge:** {r.num_tables_after_merge}")

                if r.excel_bytes:
                    base_name = os.path.splitext(r.filename)[0]
                    st.download_button(
                        label=f"⬇️ Download {base_name}.xlsx",
                        data=r.excel_bytes,
                        file_name=f"{base_name}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key=f"dl_{r.filename}",
                    )

                if r.preview_frames:
                    st.markdown("**Preview:**")
                    for sheet_name, preview_df in r.preview_frames.items():
                        st.caption(f"Sheet / View: {sheet_name}")
                        st.dataframe(preview_df, use_container_width=True)
    else:
        st.info("Upload PDF files above and click **Start Conversion** to begin.")


if __name__ == "__main__":
    main()
